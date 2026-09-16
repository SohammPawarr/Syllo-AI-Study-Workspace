"""Multi-format parsing service (PDF, DOCX, PPTX) to Markdown and semantic chunking."""

import io
import os
import requests
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
from langchain_text_splitters import MarkdownHeaderTextSplitter

def extract_markdown_from_file(file_path_or_url: str) -> str:
    """Extracts text from PDF, DOCX, or PPTX and converts it into Markdown."""
    
    if file_path_or_url.startswith(("http://", "https://")):
        response = requests.get(file_path_or_url, timeout=60)
        response.raise_for_status()
        file_bytes = response.content
        # Try to infer extension from url, defaulting to pdf
        ext = os.path.splitext(file_path_or_url.split("?")[0])[1].lower()
        if not ext:
            ext = ".pdf"
    else:
        with open(file_path_or_url, "rb") as f:
            file_bytes = f.read()
        ext = os.path.splitext(file_path_or_url)[1].lower()

    if ext == ".docx":
        return _extract_docx(file_bytes)
    elif ext == ".pptx":
        return _extract_pptx(file_bytes)
    else:
        return _extract_pdf(file_bytes)

def _extract_pdf(pdf_bytes: bytes) -> str:
    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    pages_text = []
    for i, page in enumerate(doc):
        text = page.get_text()
        if text:
            # basic structure to mimic markdown headers per page
            pages_text.append(f"## Page {i+1}\n\n{text.strip()}")
    return "\n\n".join(pages_text)

def _extract_docx(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    md_lines = []
    for para in doc.paragraphs:
        if para.style.name.startswith('Heading'):
            level = para.style.name.replace('Heading ', '')
            try:
                level = int(level)
            except:
                level = 1
            md_lines.append(f"{'#' * level} {para.text}")
        elif para.text.strip():
            md_lines.append(para.text)
    return "\n\n".join(md_lines)

def _extract_pptx(pptx_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(pptx_bytes))
    md_lines = []
    for i, slide in enumerate(prs.slides):
        md_lines.append(f"## Slide {i+1}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                md_lines.append(shape.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                md_lines.append(f"### Speaker Notes\n{notes}")
    return "\n\n".join(md_lines)

def chunk_markdown(raw_markdown: str) -> list[dict]:
    """
    Split markdown semantically by headers.
    """
    headers_to_split_on = [
        ("#", "Header 1"),
        ("##", "Header 2"),
        ("###", "Header 3"),
    ]
    markdown_splitter = MarkdownHeaderTextSplitter(headers_to_split_on=headers_to_split_on)
    splits = markdown_splitter.split_text(raw_markdown)
    
    # Format chunks to match existing structure
    return [{"index": i, "text": f"{split.metadata}\n{split.page_content}"} for i, split in enumerate(splits)]
